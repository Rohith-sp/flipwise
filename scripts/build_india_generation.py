"""Grid-India raw files → data/raw/india/india_generation.csv (PRD §6.2.2, source option 3).

Reads the files fetched by scripts/download_india.py:
  - daily PSP reports (sheet MOP_E): section G "Sourcewise generation (Gross)" gives regional
    daily MU by fuel; section A gives regional wind and solar MU, which split the RES total;
  - Mendeley SCADA workbooks: all-India MW by fuel every 10 s, averaged into 30-minute IST slots.
Then calls flipwise.ingest.construct_india_generation and writes `region_id,slot_utc,fuel,mwh`.

Needs the optional extra: uv sync --extra india-data
Usage: uv run python scripts/build_india_generation.py
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
from pathlib import Path

import pandas as pd

from flipwise.ingest import IST_OFFSET_S, construct_india_generation

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw" / "india"
REGION_ID = {"NR": 1, "WR": 2, "SR": 3, "ER": 4}  # NER unused (PRD §6.2.1)
PSP_FUEL = {
    "coal": "coal",
    "lignite": "lignite",
    "hydro": "hydro",
    "nuclear": "nuclear",
    "gas, naptha & diesel": "gas",
    "gas, naphtha & diesel": "gas",
}
SCADA_SHAPE = {
    "THM": "thermal",
    "HYD": "hydro",
    "GAS": "gas",
    "NPC": "nuclear",
    "WIND": "wind",
    "SOLAR": "solar",
}


def _rows(path: Path) -> list[list]:
    if path.suffix == ".xls":
        import xlrd

        sh = next(s for s in xlrd.open_workbook(path).sheets() if "MOP" in s.name.upper())
        return [sh.row_values(r) for r in range(sh.nrows)]
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[next(n for n in wb.sheetnames if "MOP" in n.upper())]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _num(x) -> float:
    if x is None or (isinstance(x, str) and x.strip() in ("", "-")):
        return 0.0
    return float(x)


def _label(row: list) -> str:
    return next((str(c).strip() for c in row if isinstance(c, str) and c.strip()), "")


def _region_cols(row: list) -> dict[str, int]:
    return {str(c).strip(): i for i, c in enumerate(row) if str(c).strip() in REGION_ID}


def _section(rows: list[list], title: str) -> tuple[int, dict[str, int]]:
    i = next(i for i, r in enumerate(rows) if _label(r).startswith(title))
    h = next(j for j in range(i + 1, i + 4) if len(_region_cols(rows[j])) == len(REGION_ID))
    return h, _region_cols(rows[h])


def parse_psp(path: Path, day: dt.date) -> list[tuple]:
    """Regional daily MWh by fuel. RES is split into wind, solar and a biomass/other rest."""
    rows = _rows(path)
    out: dict[tuple[str, str], float] = {}
    h, cols = _section(rows, "G. Sourcewise generation")
    res: dict[str, float] = {}
    for r in rows[h + 1 : h + 12]:
        lab = _label(r).lower()
        if lab.startswith("total"):
            break
        if lab.startswith("res"):
            res = {reg: _num(r[c]) for reg, c in cols.items()}
        elif lab in PSP_FUEL:
            for reg, c in cols.items():
                out[(reg, PSP_FUEL[lab])] = out.get((reg, PSP_FUEL[lab]), 0.0) + _num(r[c])
    h, cols = _section(rows, "A. Power Supply Position")
    for r in rows[h + 1 : h + 12]:
        lab = _label(r).lower()
        for key, fuel in (("wind gen", "wind"), ("solar gen", "solar")):
            if lab.startswith(key):
                for reg, c in cols.items():
                    out[(reg, fuel)] = _num(r[c])
    for reg in REGION_ID:
        if reg not in res:
            raise ValueError(f"{path.name}: no RES row in section G")
        out[(reg, "biomass")] = max(
            0.0, res[reg] - out.get((reg, "wind"), 0.0) - out.get((reg, "solar"), 0.0)
        )
    return [(day, REGION_ID[reg], fuel, mu * 1000.0) for (reg, fuel), mu in sorted(out.items())]


def _to_long(df: pd.DataFrame) -> pd.DataFrame:
    """`ist` + one column per shape → `slot_utc, shape, mw` (mean per 30-minute slot)."""
    utc_s = df["ist"].astype("datetime64[s]").astype("int64") - IST_OFFSET_S
    df = df.drop(columns="ist").assign(slot_utc=(utc_s // 1800) * 1800)
    long = df.melt(id_vars="slot_utc", var_name="shape", value_name="mw")
    long["mw"] = pd.to_numeric(long["mw"], errors="coerce")
    return long.dropna().groupby(["slot_utc", "shape"], as_index=False)["mw"].mean()


def _parse_hourly_report(ws) -> pd.DataFrame:
    """The 2024–25 workbook: hourly 'Timestamp, Demand, Wind, Solar, Total' only.

    Each hourly value stands for both half-hours of that hour. Thermal, hydro, gas and nuclear
    are absent, so those days take a typical shape in construct_india_generation.
    """
    it = ws.iter_rows(values_only=True)
    head = [str(c or "").lower() for c in next(it)]
    iw, isol = head.index("wind (mw)"), head.index("solar (mw)")
    recs = []
    for r in it:
        try:
            t = dt.datetime.strptime(str(r[0]), "%d-%m-%Y %H:%M:%S")
        except ValueError:
            continue  # the trailing "Sum" row
        for half in (0, 30):
            recs.append((t + dt.timedelta(minutes=half), r[iw], r[isol]))
    return _to_long(pd.DataFrame(recs, columns=["ist", "wind", "solar"]))


def parse_scada(path: Path) -> pd.DataFrame:
    """All-India MW by fuel → mean per 30-minute slot. SCADA timestamps are IST."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Report" in wb.sheetnames:
        return _parse_hourly_report(wb["Report"])
    ws = wb[next(n for n in wb.sheetnames if n.lower().startswith("sheet1"))]
    it = ws.iter_rows(values_only=True)
    next(it)
    tags = next(it)
    cols = {}
    for i, t in enumerate(tags):
        tag = str(t or "").upper()
        if "|P" not in tag:
            continue
        for key, shape in SCADA_SHAPE.items():
            if key in tag and shape not in cols and "DEMAND" not in tag:
                cols[shape] = i
    if len(cols) != len(SCADA_SHAPE):
        raise ValueError(f"{path.name}: SCADA columns {sorted(cols)} incomplete")
    recs = [(r[0], *(r[i] for i in cols.values())) for r in it if isinstance(r[0], dt.datetime)]
    return _to_long(pd.DataFrame(recs, columns=["ist", *cols.keys()]))


def fill_missing_days(daily: pd.DataFrame) -> tuple[pd.DataFrame, list[dt.date]]:
    """A day without a report takes the mean of the nearest reported days before and after.

    A whole missing day would otherwise be a 48-slot truth gap, which fails `validate`
    (PRD §6.2.2 allows interpolation over at most 2 slots). Filled days are listed in the
    manifest. Runs of more than 2 missing days are left as gaps on purpose.
    """
    days = sorted(set(daily["ist_date"]))
    filled, extra = [], []
    for a, b in zip(days, days[1:], strict=False):
        gap = (b - a).days - 1
        if not 1 <= gap <= 2:
            continue
        both = pd.concat([daily[daily["ist_date"] == a], daily[daily["ist_date"] == b]])
        mean = both.groupby(["region_id", "fuel"], as_index=False)["mwh_day"].mean()
        for i in range(1, gap + 1):
            d = a + dt.timedelta(days=i)
            extra.append(mean.assign(ist_date=d))
            filled.append(d)
    if extra:
        daily = pd.concat([daily, *extra], ignore_index=True)
    return daily, filled


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", type=Path, default=RAW / "india_generation.csv")
    args = ap.parse_args()

    daily, bad = [], []
    for f in sorted((RAW / "grid-india" / "psp").glob("*_NLDC_PSP.xls*")):
        day = dt.date.fromisoformat(f.name[:10])
        try:
            daily += parse_psp(f, day)
        except (StopIteration, ValueError) as e:
            bad.append(f"{f.name}: {e or 'section not found'}")
    daily_df = pd.DataFrame(daily, columns=["ist_date", "region_id", "fuel", "mwh_day"])
    print(f"PSP: {daily_df['ist_date'].nunique()} days parsed, {len(bad)} unreadable")
    daily_df, filled = fill_missing_days(daily_df)
    if filled:
        print(f"PSP: filled {len(filled)} missing days from their neighbours: {filled}")

    profiles = []
    for f in sorted((RAW / "mendeley").glob("*.xlsx")):
        p = parse_scada(f)
        print(f"SCADA {f.name}: {len(p)} slot-values")
        profiles.append(p)
    profile = pd.concat(profiles, ignore_index=True)
    profile = profile.groupby(["slot_utc", "shape"], as_index=False)["mw"].mean()

    gen, counts = construct_india_generation(daily_df, profile)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    gen.to_csv(args.out, index=False, lineterminator="\n")
    report = {
        "psp_days": int(daily_df["ist_date"].nunique()),
        "psp_unreadable": bad,
        "psp_days_filled_from_neighbours": [d.isoformat() for d in filled],
        "scada_slots": int(profile["slot_utc"].nunique()),
        **counts,
        "rows": len(gen),
        "output": args.out.relative_to(ROOT).as_posix(),
    }
    (ROOT / "data" / "manifests" / "india_generation.json").write_text(
        json.dumps(report, indent=2, default=str) + "\n"
    )
    print(json.dumps(report, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
